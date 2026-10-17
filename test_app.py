import os
import io
import app

def test_full_pipeline():
    print("Iniciando testes integrados da aplicacao SplitC...")
    client = app.app.test_client()

    # 1. Teste da rota inicial e favicon
    res = client.get('/')
    assert res.status_code == 200, f"Falha GET /: {res.status_code}"
    print("[OK] GET / funcionando perfeitamente (HTTP 200)")

    res_fav = client.get('/favicon.ico')
    assert res_fav.status_code == 200, f"Falha GET /favicon.ico: {res_fav.status_code}"
    assert len(res_fav.data) > 0
    print("[OK] GET /favicon.ico funcionando (HTTP 200, logo carregado)")

    # 2. Teste de Upload dos arquivos de teste
    with open('planilha_antiga.xlsx', 'rb') as fa, open('planilha_nova.xlsx', 'rb') as fb:
        data = {
            'file_a': (io.BytesIO(fa.read()), 'planilha_antiga.xlsx'),
            'file_b': (io.BytesIO(fb.read()), 'planilha_nova.xlsx')
        }
        res = client.post('/api/upload-headers', data=data, content_type='multipart/form-data')
    
    assert res.status_code == 200, f"Falha POST /api/upload-headers: {res.status_code}"
    upload_res = res.get_json()
    assert 'session_id' in upload_res
    session_id = upload_res['session_id']
    print(f"[OK] Upload e leitura de cabecalhos realizado. Sessao: {session_id}")
    print(f"  Colunas detectadas em A: {len(upload_res['cols_a'])}")
    print(f"  Colunas detectadas em B: {len(upload_res['cols_b'])}")
    print(f"  Mapeamento de Data_Fechamento: {upload_res['column_mapping'].get('Data_Fechamento')}")

    # 3. Teste com seletor único da esquerda (sem colunas complementares)
    compare_payload_single = {
        'session_id': session_id,
        'primary_key_a': 'SEQ_COM_SUM_VND',
        'primary_key_b': 'SEQ_COM_SUM_VND',
        'secondary_keys_a': [],
        'normalize_currency': True,
        'normalize_dates': True,
        'ignore_case': True,
        'trim_spaces': True,
        'column_mapping': upload_res['column_mapping']
    }
    res = client.post('/api/compare', json=compare_payload_single)
    assert res.status_code == 200
    comp_single = res.get_json()
    assert comp_single['summary']['is_identical'] is True
    print("[OK] Teste com apenas a chave principal da esquerda: 100% identico")

    # 4. Teste com 1 coluna complementar na direita (ex: DATA_VENDA)
    compare_payload_one_sec = dict(compare_payload_single)
    compare_payload_one_sec['secondary_keys_a'] = ['DATA_VENDA']
    res = client.post('/api/compare', json=compare_payload_one_sec)
    assert res.status_code == 200
    comp_one = res.get_json()
    assert comp_one['summary']['is_identical'] is True
    print("[OK] Teste com 1 coluna complementar (DATA_VENDA): 100% identico")

    # 5. Teste com 2 ou mais colunas complementares na direita (ex: DATA_VENDA e FILIAL)
    compare_payload_multi_sec = dict(compare_payload_single)
    compare_payload_multi_sec['secondary_keys_a'] = ['DATA_VENDA', 'FILIAL']
    res = client.post('/api/compare', json=compare_payload_multi_sec)
    assert res.status_code == 200
    comp_multi = res.get_json()
    assert comp_multi['summary']['is_identical'] is True
    assert comp_multi['summary']['secondary_keys_a'] == ['DATA_VENDA', 'FILIAL']
    print("[OK] Teste com 2 colunas complementares (DATA_VENDA + FILIAL): 100% identico")

    # 6. Teste com divergencias intencionais (desativando normalizacao de moedas)
    compare_payload_diff = dict(compare_payload_multi_sec)
    compare_payload_diff['normalize_currency'] = False

    res = client.post('/api/compare', json=compare_payload_diff)
    assert res.status_code == 200
    comp_res_diff = res.get_json()
    print("[OK] Comparacao sem normalizacao de moeda (cenario com divergencias):")
    print(f"  Bases identicas? {comp_res_diff['summary']['is_identical']}")
    print(f"  Linhas divergentes: {comp_res_diff['summary']['different_rows_count']}")
    print(f"  Celulas divergentes: {comp_res_diff['summary']['total_cells_diff']}")
    assert comp_res_diff['summary']['is_identical'] is False
    assert len(comp_res_diff['diff_preview']) > 0

    # 7. Teste do download do relatorio Excel gerado com as celulas pintadas
    report_id = comp_res_diff['report_id']
    res_down = client.get(f'/api/download-report/{report_id}')
    assert res_down.status_code == 200
    assert len(res_down.data) > 10000, "Relatorio Excel gerado esta vazio ou corrompido!"

    # Valida que o Excel gerado contém as abas esperadas, incluindo a nova 3ª aba 'Sumarização Financeira'
    import openpyxl
    wb_test = openpyxl.load_workbook(io.BytesIO(res_down.data))
    assert 'Divergências Detalhadas' in wb_test.sheetnames
    assert 'Sumarização Financeira' in wb_test.sheetnames, "Aba 'Sumarização Financeira' deve existir quando houver divergências!"
    
    # Valida presença da logo no Resumo
    ws_resumo = wb_test['Resumo da Validação']
    assert len(ws_resumo._images) > 0, "Logo SplitC deve estar presente lateralmente na aba de resumo!"
    print(f"[OK] Logo da SplitC confirmada no Excel ({len(ws_resumo._images)} imagem inserida).")

    # Valida cabeçalhos e cores (#4FD580)
    ws_test = wb_test['Divergências Detalhadas']
    headers_test = [ws_test.cell(row=1, column=i).value for i in range(1, 6)]
    expected_headers = ["Chave Identificadora", "Linha", "Coluna", "Valor Base Antiga", "Valor Base Nova"]
    assert headers_test == expected_headers, f"Headers incorretos: {headers_test} vs {expected_headers}"
    assert ws_test.cell(row=2, column=2).value is not None, "Coluna Linha vazia na linha 2"
    assert ws_test.cell(row=1, column=1).fill.start_color.rgb == '004FD580', "Cabeçalho deve usar a cor #4FD580!"

    # Valida a 3ª aba 'Sumarização Financeira'
    ws_fin = wb_test['Sumarização Financeira']
    fin_headers = [ws_fin.cell(row=1, column=i).value for i in range(1, 9)]
    expected_fin_headers = [
        "Coluna Base Antiga", "Coluna Base Nova", "Total Base Antiga (R$)", "Total Base Nova (R$)",
        "Diferença (Nova - Antiga)", "% Variação", "Situação", "Qtd. Linhas Divergentes"
    ]
    assert fin_headers == expected_fin_headers, f"Cabeçalhos da aba financeira incorretos: {fin_headers}"
    assert ws_fin.cell(row=1, column=1).fill.start_color.rgb == '004FD580', "Cabeçalhos da aba financeira devem usar #4FD580!"
    assert ws_fin.cell(row=2, column=3).value is not None, "Total da Base Antiga deve estar preenchido!"
    print(f"[OK] 3ª Aba 'Sumarização Financeira' confirmada com cabeçalhos em #4FD580: {fin_headers}")

    # Valida que em bases idênticas a 3ª aba 'Sumarização Financeira' NÃO é criada
    report_id_single = comp_single['report_id']
    res_down_single = client.get(f'/api/download-report/{report_id_single}')
    wb_single = openpyxl.load_workbook(io.BytesIO(res_down_single.data))
    assert 'Sumarização Financeira' not in wb_single.sheetnames, "Aba financeira NÃO deve existir em bases 100% idênticas!"
    print("[OK] Verificado: em bases 100% idênticas, a terceira aba não é criada (mantida estrutura enxuta).")

    # 8. Teste de Mapeamento Manual de Colunas com nomes completamente diferentes (+ Adicionar Coluna)
    # Simula o usuário adicionando manualmente correspondências que o fuzzy matching não detectou
    import pandas as pd
    df_custom_a = pd.DataFrame({
        'MATRICULA_REF': ['M001', 'M002'],
        'SALARIO_BASE': [5000.0, 7500.0],
        'SETOR_ATUACAO': ['Financeiro', 'TI']
    })
    df_custom_b = pd.DataFrame({
        'COD_FUNCIONARIO': ['M001', 'M002'],
        'REMUNERACAO_BRUTA': ['5.000,00', '7.500,00'],
        'DEPARTAMENTO': ['Financeiro', 'TI']
    })
    buf_a = io.BytesIO()
    buf_b = io.BytesIO()
    with pd.ExcelWriter(buf_a, engine='openpyxl') as writer:
        df_custom_a.to_excel(writer, index=False)
    with pd.ExcelWriter(buf_b, engine='openpyxl') as writer:
        df_custom_b.to_excel(writer, index=False)
    buf_a.seek(0)
    buf_b.seek(0)

    data_custom = {
        'file_a': (buf_a, 'base_antiga_custom.xlsx'),
        'file_b': (buf_b, 'base_nova_custom.xlsx')
    }
    res_upload_custom = client.post('/api/upload-headers', data=data_custom, content_type='multipart/form-data')
    assert res_upload_custom.status_code == 200
    custom_session = res_upload_custom.get_json()

    # O fuzzy matching não mapeia porque os nomes são 100% diferentes
    assert len(custom_session['column_mapping']) == 0

    # Simula o mapeamento manual pelo botão '+' do front-end
    manual_custom_mapping = {
        'MATRICULA_REF': 'COD_FUNCIONARIO',
        'SALARIO_BASE': 'REMUNERACAO_BRUTA',
        'SETOR_ATUACAO': 'DEPARTAMENTO'
    }

    res_compare_custom = client.post('/api/compare', json={
        'session_id': custom_session['session_id'],
        'primary_key_a': 'MATRICULA_REF',
        'primary_key_b': 'COD_FUNCIONARIO',
        'secondary_keys_a': [],
        'normalize_currency': True,
        'normalize_dates': True,
        'ignore_case': True,
        'trim_spaces': True,
        'column_mapping': manual_custom_mapping
    })
    assert res_compare_custom.status_code == 200
    custom_result = res_compare_custom.get_json()
    assert custom_result['summary']['is_identical'] is True
    assert custom_result['summary']['total_columns_compared'] == 3
    # 9. Teste da feature de Limpar Mapeamentos e selecionar apenas colunas desejadas
    # Simula o usuário clicando no botão vermelho de Limpar Mapeamentos e escolhendo apenas 1 coluna
    res_compare_cleared = client.post('/api/compare', json={
        'session_id': session_id,
        'primary_key_a': 'SEQ_COM_SUM_VND',
        'primary_key_b': 'SEQ_COM_SUM_VND',
        'secondary_keys_a': [],
        'normalize_currency': True,
        'normalize_dates': True,
        'ignore_case': True,
        'trim_spaces': True,
        'column_mapping': {'Data_Fechamento': 'Data Fechametno'}  # Apenas 1 coluna selecionada pelo usuário
    })
    assert res_compare_cleared.status_code == 200
    cleared_result = res_compare_cleared.get_json()
    assert cleared_result['summary']['total_columns_compared'] == 1, f"Esperado 1 coluna comparada, obtido: {cleared_result['summary']['total_columns_compared']}"
    print("[OK] Teste de Limpar Mapeamento e comparar apenas colunas selecionadas: 1 coluna comparada com sucesso!")

    # 10. Teste da página Validador de Texto (/validador-texto) e rotas de comparação (.txt, .docx, .pdf) + Exportação PDF
    res_txt_page = client.get('/validador-texto')
    assert res_txt_page.status_code == 200, f"Falha GET /validador-texto: {res_txt_page.status_code}"
    html_txt_page = res_txt_page.data.decode('utf-8')
    assert 'Rich text' in html_txt_page
    assert 'Redline' in html_txt_page
    assert 'Plain text' in html_txt_page
    assert 'OCR' not in html_txt_page
    print("[OK] GET /validador-texto renderizado com os 3 modos (Rich text, Redline, Plain text) e sem Imagem/OCR!")

    # Teste 10.1: Upload de arquivos .docx reais
    import docx
    doc_a_buf = io.BytesIO()
    doc_a = docx.Document()
    doc_a.add_paragraph("Contrato de Prestação de Serviços ValidAí")
    doc_a.add_paragraph("Cláusula 1: O valor mensal será de R$ 3.000,00.")
    doc_a.add_paragraph("Cláusula 2: Vigência de 12 meses.")
    doc_a.save(doc_a_buf)
    doc_a_buf.seek(0)

    doc_b_buf = io.BytesIO()
    doc_b = docx.Document()
    doc_b.add_paragraph("Contrato de Prestação de Serviços ValidAí")
    doc_b.add_paragraph("Cláusula 1: O valor mensal será de R$ 4.500,00.")
    doc_b.add_paragraph("Cláusula 3: Foro da comarca de São Paulo.")
    doc_b.save(doc_b_buf)
    doc_b_buf.seek(0)

    res_docx_cmp = client.post(
        '/api/text/compare',
        data={
            'file_a': (doc_a_buf, 'contrato_antigo.docx'),
            'file_b': (doc_b_buf, 'contrato_novo.docx'),
            'trim_spaces': 'true',
            'ignore_case': 'false',
            'diff_granularity': 'word'
        },
        content_type='multipart/form-data'
    )
    assert res_docx_cmp.status_code == 200, f"Falha POST /api/text/compare (.docx): {res_docx_cmp.status_code}"
    docx_json = res_docx_cmp.get_json()
    assert docx_json['summary']['is_identical'] is False
    assert docx_json['summary']['file_type_a'] == 'DOCX'
    assert docx_json['summary']['total_diff_blocks'] >= 2
    assert 'report_id' in docx_json

    # Verifica download do relatório PDF gerado
    pdf_report_id = docx_json['report_id']
    res_pdf_down = client.get(f'/api/text/download-pdf/{pdf_report_id}')
    assert res_pdf_down.status_code == 200, f"Falha download PDF: {res_pdf_down.status_code}"
    assert res_pdf_down.data.startswith(b'%PDF-'), "Arquivo baixado não é um PDF válido!"
    assert len(res_pdf_down.data) > 1500
    print(f"[OK] Comparação de .docx e exportação de relatório PDF validada ({len(res_pdf_down.data)} bytes)!")

    # Teste 10.2: Upload de arquivos .pdf reais (usando o próprio PDF gerado ou criando um PDF de teste)
    pdf_a_buf = io.BytesIO(res_pdf_down.data)
    pdf_b_buf = io.BytesIO(res_pdf_down.data)
    res_pdf_cmp = client.post(
        '/api/text/compare',
        data={
            'file_a': (pdf_a_buf, 'relatorio_v1.pdf'),
            'file_b': (pdf_b_buf, 'relatorio_v2.pdf'),
        },
        content_type='multipart/form-data'
    )
    assert res_pdf_cmp.status_code == 200, f"Falha POST /api/text/compare (.pdf): {res_pdf_cmp.status_code}"
    pdf_cmp_json = res_pdf_cmp.get_json()
    assert pdf_cmp_json['summary']['file_type_a'] == 'PDF'
    assert pdf_cmp_json['summary']['is_identical'] is True
    print("[OK] Leitura e comparação de arquivos .pdf funcionando perfeitamente!")

    # 11. Teste da página "Como usar?" (/como-usar) com Toggle Planilhas / Texto
    res_guide = client.get('/como-usar')
    assert res_guide.status_code == 200, f"Falha GET /como-usar: {res_guide.status_code}"
    html_guide = res_guide.data.decode('utf-8')
    assert 'btn-guide-planilhas' in html_guide
    assert 'btn-guide-texto' in html_guide
    assert 'guide-view-planilhas' in html_guide
    assert 'guide-view-texto' in html_guide
    print("[OK] GET /como-usar renderizado com toggle Validar Planilhas / Validar Texto e ambos os guias!")

    # 12. Teste do Dropdown de Idiomas (PT, ES, EN) e Internacionalização (i18n) em todas as páginas
    res_i18n_js = client.get('/static/js/i18n.js')
    assert res_i18n_js.status_code == 200, f"Falha GET /static/js/i18n.js: {res_i18n_js.status_code}"
    i18n_content = res_i18n_js.data.decode('utf-8')
    assert 'pt:' in i18n_content, "Dicionário PT ausente em i18n.js"
    assert 'es:' in i18n_content, "Dicionário ES ausente em i18n.js"
    assert 'en:' in i18n_content, "Dicionário EN ausente em i18n.js"
    print("[OK] Arquivo estático i18n.js carregado com sucesso contendo os 3 idiomas (PT, ES, EN)!")

    # Checar presença do dropdown e scripts em /, /validador-texto e /como-usar
    pages_to_check = [
        ('/', 'Validador Planilhas'),
        ('/validador-texto', 'Validador Texto'),
        ('/validador-cadastros', 'Higienizador de Cadastros'),
        ('/como-usar', 'Como Usar')
    ]
    for path, label in pages_to_check:
        res_p = client.get(path)
        assert res_p.status_code == 200, f"Falha GET {path}: {res_p.status_code}"
        html_p = res_p.data.decode('utf-8')
        assert 'lang-dropdown-wrapper' in html_p, f"Dropdown de idiomas ausente na página {label} ({path})"
        assert 'lang-toggle-btn' in html_p, f"Botão de alternância de idioma ausente na página {label} ({path})"
        assert 'data-lang="pt"' in html_p, f"Opção Português ausente na página {label} ({path})"
        assert 'data-lang="es"' in html_p, f"Opção Espanhol ausente na página {label} ({path})"
        assert 'data-lang="en"' in html_p, f"Opção Inglês ausente na página {label} ({path})"
        assert 'i18n.js' in html_p, f"Script i18n.js não importado na página {label} ({path})"
        assert 'data-i18n=' in html_p, f"Atributos data-i18n ausentes na página {label} ({path})"
        assert 'data-i18n="footer_copyright"' in html_p, f"Footer com data-i18n ausente na página {label} ({path})"
        print(f"[OK] Página {label} ({path}) validada com dropdown de idiomas e i18n!")

    # 13. Teste do Menu Dropdown de Soluções ("Mais Validadores")
    for path, label in pages_to_check:
        res_p = client.get(path)
        assert res_p.status_code == 200
        html_p = res_p.data.decode('utf-8')
        assert 'nav-solutions-dropdown' in html_p, f"Dropdown Mais Validadores ausente em {label}"
        assert 'nav-solutions-btn' in html_p, f"Botão Mais Validadores ausente em {label}"
        assert 'nav-solutions-menu' in html_p, f"Menu Mais Validadores ausente em {label}"
        assert '/validador-cadastros' in html_p, f"Item Higienizador de Cadastros ausente em {label}"
        assert '/validador-xml' in html_p, f"Item Conciliador XML ausente em {label}"
        assert '/conciliador-taxas' in html_p, f"Item Conciliação Taxas OFX ausente em {label}"
        assert '/auditoria-rh' in html_p, f"Item Auditoria RH ausente em {label}"
        assert '/validador-json' in html_p, f"Item Comparador JSON ausente em {label}"
        assert 'badge_coming_soon' in html_p, f"Badge 'Em breve' ausente em {label}"
        assert 'Em breve' in html_p, f"Texto da badge 'Em breve' ausente em {label}"
    print("[OK] Menu Dropdown 'Mais Validadores' validado com os 5 novos módulos em todas as páginas!")

    # 14. Teste completo do Higienizador de Cadastros (/validador-cadastros)
    # A. Renderização da página e seus elementos de UI
    res_cad = client.get('/validador-cadastros')
    assert res_cad.status_code == 200, f"Falha GET /validador-cadastros: {res_cad.status_code}"
    html_cad = res_cad.data.decode('utf-8')
    assert 'drop-cadastros' in html_cad, "Área de upload único ausente"
    assert 'step-config-cadastros' in html_cad, "Card de configurações e regras ausente"
    assert 'step-results-cadastros' in html_cad, "Seção de resultados ausente"
    assert 'kpi-total-rows' in html_cad, "KPI total analisado ausente"
    assert 'kpi-valid-rows' in html_cad, "KPI válidos ausente"
    assert 'kpi-invalid-docs' in html_cad, "KPI docs inválidos ausente"
    assert 'kpi-duplicate-rows' in html_cad, "KPI duplicados ausente"
    assert 'kpi-suspicious-contacts' in html_cad, "KPI contatos suspeitos ausente"
    assert 'kpi-auto-corrections' in html_cad, "KPI auto correções ausente"
    assert 'btn-download-clean' in html_cad, "Botão download base limpa ausente"
    assert 'btn-download-audit' in html_cad, "Botão download auditoria ausente"
    print("[OK] Página /validador-cadastros renderizada com todos os elementos de UI!")

    # B. Teste de Upload de Base de Cadastros e Leitura de Cabeçalhos
    import pandas as pd
    test_cadastros_df = pd.DataFrame({
        'CPF_Cliente': [
            '52998224725',        # 1: CPF Válido
            '1234567890',         # 2: 10 dígitos (recupera zero à esquerda para '01234567890')
            '529.982.247-25',     # 3: Duplicado da linha 1
            '111.111.111-11',     # 4: CPF Inválido (dígitos repetidos)
            '00.000.000/0001-91'  # 5: CNPJ Válido (Banco do Brasil)
        ],
        'Nome_Razao': [
            'MARIA DA SILVA',
            'joão  souza 123',
            'Maria da Silva',
            'Cliente Invalido ---',
            'BANCO DO BRASIL S/A'
        ],
        'Email_Contato': [
            'MARIA@GMAIL.COM',
            'joao@gmial.com',     # Domínio com typo conhecido (@gmial.com)
            'maria@gmail.com',
            'carlos_invalido',    # Sem arroba
            'contato@bb.com.br'
        ],
        'Telefone_Celular': [
            '11987654321',        # Celular SP válido
            '1187654321',         # Fixo SP válido (8 dígitos)
            '11 98765-4321',      # Formatado
            '001234',             # DDD inválido
            '61 3493-9000'        # Fixo Brasília válido
        ],
        'Cidade': ['São Paulo', 'Campinas', 'São Paulo', 'Santos', 'Brasília']
    })
    
    excel_buf = io.BytesIO()
    test_cadastros_df.to_excel(excel_buf, index=False)
    excel_buf.seek(0)

    res_upload_cad = client.post(
        '/api/cadastros/upload-headers',
        data={'file': (excel_buf, 'clientes_teste.xlsx')},
        content_type='multipart/form-data'
    )
    assert res_upload_cad.status_code == 200, f"Falha POST /api/cadastros/upload-headers: {res_upload_cad.status_code}"
    cad_upload_json = res_upload_cad.get_json()
    assert 'session_id' in cad_upload_json
    session_id_cad = cad_upload_json['session_id']
    assert cad_upload_json['total_rows'] == 5
    assert len(cad_upload_json['cols']) == 5
    mapping = cad_upload_json['suggested']
    assert mapping['col_doc'] == 'CPF_Cliente'
    assert mapping['col_name'] == 'Nome_Razao'
    assert mapping['col_email'] == 'Email_Contato'
    assert mapping['col_phone'] == 'Telefone_Celular'
    print("[OK] Upload de base de cadastros e detecção automática de colunas realizado com sucesso!")

    # C. Teste do Processamento e Regras de Higienização
    process_payload = {
        'session_id': session_id_cad,
        'col_doc': 'CPF_Cliente',
        'col_name': 'Nome_Razao',
        'col_email': 'Email_Contato',
        'col_phone': 'Telefone_Celular',
        'duplicate_keys': ['doc'],
        'doc_format': 'mask',
        'auto_recover_zero': True,
        'name_title_case': True,
        'email_lowercase': True,
        'phone_format': 'mask',
        'remove_duplicates': False
    }
    res_proc = client.post('/api/cadastros/process', json=process_payload)
    assert res_proc.status_code == 200, f"Falha POST /api/cadastros/process: {res_proc.status_code}"
    proc_json = res_proc.get_json()
    assert proc_json['success'] is True
    summary = proc_json['summary']
    assert summary['total_rows'] == 5
    assert summary['invalid_docs_count'] >= 1, "Deveria detectar CPF inválido"
    assert summary['duplicates_count'] >= 1, "Deveria detectar CPF duplicado"
    assert summary['suspicious_contacts_count'] >= 1, "Deveria detectar typo no email e telefone com erro"
    assert summary['auto_corrections_count'] >= 1, "Deveria registrar correções automáticas de zeros/nomes/emails"
    print(f"[OK] Processamento de cadastros realizado: {summary['valid_count']} válidos, {summary['invalid_docs_count']} inválidos, {summary['duplicates_count']} duplicados, {summary['auto_corrections_count']} correções automáticas.")

    # D. Teste de Download da Base 100% Higienizada (Excel)
    report_id_cad = proc_json['report_id']
    res_down_clean = client.get(f'/api/cadastros/download-clean/{report_id_cad}')
    assert res_down_clean.status_code == 200, f"Falha GET download-clean: {res_down_clean.status_code}"
    assert len(res_down_clean.data) > 3000, "Base limpa gerada está vazia"
    wb_clean = openpyxl.load_workbook(io.BytesIO(res_down_clean.data))
    assert 'Base Higienizada' in wb_clean.sheetnames
    print("[OK] Download da Base 100% Higienizada gerado e validado com sucesso!")

    # E. Teste de Download do Relatório de Auditoria com Células Pintadas e Logo
    res_down_audit = client.get(f'/api/cadastros/download-audit/{report_id_cad}')
    assert res_down_audit.status_code == 200, f"Falha GET download-audit: {res_down_audit.status_code}"
    assert len(res_down_audit.data) > 5000, "Relatório de auditoria gerado está vazio"
    wb_audit = openpyxl.load_workbook(io.BytesIO(res_down_audit.data))
    assert 'Resumo da Auditoria' in wb_audit.sheetnames
    assert any('Base Auditada' in s for s in wb_audit.sheetnames)
    assert any('Inconsist' in s for s in wb_audit.sheetnames)
    ws_aud_resumo = wb_audit['Resumo da Auditoria']
    assert len(ws_aud_resumo._images) >= 1, "Logo ValidAí ausente no Resumo da Auditoria!"
    print("[OK] Relatório de Auditoria gerado com células pintadas, 3 abas e Logo ValidAí!")

    # F. Teste de Coluna Incompatível (ex: IDCOLABORADOR com números aleatórios curtos como '555722')
    import pandas as pd
    df_incompat = pd.DataFrame({
        'IDCOLABORADOR': [555722, 10293, 88392, 44021, 99120],  # 5-6 dígitos (IDs numéricos com dtype int64)
        'NMCLIENTE': ['MARIA SILVA', 'JOAO SOUZA', 'ANA COSTA', 'PEDRO LIMA', 'LUCAS DIAS'],
        'EMAIL': ['m@test.com', 'j@test.com', 'a@test.com', 'p@test.com', 'l@test.com']
    })
    buf_incompat = io.BytesIO()
    df_incompat.to_csv(buf_incompat, index=False)
    buf_incompat.seek(0)
    res_up_inc = client.post(
        '/api/cadastros/upload-headers',
        data={'file': (buf_incompat, 'base_com_ids.csv')},
        content_type='multipart/form-data'
    )
    assert res_up_inc.status_code == 200
    inc_json = res_up_inc.get_json()
    sess_inc = inc_json['session_id']

    # Tenta processar passando a coluna IDCOLABORADOR como CPF/CNPJ
    res_proc_inc = client.post('/api/cadastros/process', json={
        'session_id': sess_inc,
        'col_doc': 'IDCOLABORADOR',
        'col_name': 'NMCLIENTE',
        'col_email': 'EMAIL',
        'duplicate_keys': ['doc']
    })
    assert res_proc_inc.status_code == 400, f"Deveria rejeitar com 400, mas retornou {res_proc_inc.status_code}"
    err_data = res_proc_inc.get_json()
    assert 'não bate com a referência de CPF / CNPJ' in err_data['error']
    assert 'IDCOLABORADOR' in err_data['error']
    print(f"[OK] Detecção de coluna incompatível ('IDCOLABORADOR' com valores de 6 dígitos) validada: {err_data['error'][:80]}...")

    print("\nTODOS OS TESTES PASSARAM COM 100% DE SUCESSO!")

if __name__ == '__main__':
    test_full_pipeline()
